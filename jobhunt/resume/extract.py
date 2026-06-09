"""Extract plain text from uploaded resume files (PDF, DOCX, MD, TXT)."""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class UnsupportedFormatError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Return plain text from a resume file's raw bytes."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _from_pdf(data)
    if ext == ".docx":
        return _from_docx(data)
    if ext in {".md", ".markdown", ".txt"}:
        return data.decode("utf-8", errors="replace")
    raise UnsupportedFormatError(
        f"Unsupported resume format {ext!r}; supported: "
        + ", ".join(sorted(SUPPORTED_EXTENSIONS))
    )


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
